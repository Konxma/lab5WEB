"""Replace the figures in the user's DOCX with verified project evidence."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\vadim\Downloads\Практична_робота_5_Ткаченко_ВМ.docx")
OUTPUT = ROOT / "report_docx" / "Практична_робота_5_Ткаченко_ВМ.docx"
FIGURES = ROOT / "report_docx" / "dark_figures"

doc = Document(SOURCE)
assert len(doc.inline_shapes) == 16, "Unexpected figure count in the user's report"

title = doc.paragraphs[13]
title.text = "ПРАКТИЧНА РОБОТА №5"
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

caption_replacements = {
    2: "Рис. 2 – Структура файлів проєкту.",
    8: "Рис. 8 – Створення гаджета з автоматичною генерацією slug (201 Created).",
}

for number in range(1, 17):
    image_paragraph = doc.paragraphs[29 + number * 2]
    assert len(image_paragraph._element.xpath(".//a:blip")) == 1
    for run in list(image_paragraph.runs):
        image_paragraph._p.remove(run._element)
    picture = FIGURES / f"fig-{number:02}.png"
    image_paragraph.add_run().add_picture(str(picture), width=Inches(6.25))
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_before = Pt(2)
    image_paragraph.paragraph_format.space_after = Pt(2)

    caption = doc.paragraphs[30 + number * 2]
    if number in caption_replacements:
        caption.text = caption_replacements[number]
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_after = Pt(4)

for run in doc.paragraphs[64].runs:
    run.text = run.text.replace("створення товарів", "створення гаджетів")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
